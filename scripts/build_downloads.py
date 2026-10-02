from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import letter
from reportlab.lib.colors import HexColor, Color
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen.canvas import Canvas

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'public' / 'downloads'
OUT.mkdir(parents=True, exist_ok=True)

terms = [
    ('Protected Health Information PHI', 'Individually identifiable information that relates to a person’s health, health care, or payment for health care.'),
    ('Electronic Protected Health Information ePHI', 'PHI that is created, received, maintained, or transmitted electronically.'),
    ('Privacy Rule', 'Sets standards for how PHI may be used and disclosed and establishes patient privacy rights.'),
    ('Security Rule', 'Requires administrative, physical, and technical safeguards to protect ePHI confidentiality, integrity, and availability.'),
    ('Breach Notification Rule', 'Requires notices after certain breaches of unsecured PHI. Workforce members report suspected incidents; designated teams assess and notify.'),
    ('Covered Entity and Business Associate', 'Covered entities include health plans, clearinghouses, and qualifying providers. Business associates perform services involving PHI for a covered entity.'),
    ('Minimum Necessary', 'When it applies, use, disclose, or request only the PHI needed for the purpose. It generally does not apply to provider-to-provider treatment disclosures or requests.'),
    ('Patient Rights', 'Individuals may have rights to access records, request amendment, receive an accounting of disclosures, and request confidential communications.'),
    ('Workforce Responsibilities', 'Use approved systems, protect passwords, avoid phishing, safeguard your workspace, access only what your job requires, and report suspected incidents promptly.'),
]

def make_docx():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8); section.bottom_margin = Inches(0.8)
    title = doc.add_paragraph(style='Title')
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run('Certificate of Completion').font.color.rgb = RGBColor(0, 0, 0)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('RevExpert One Self-Paced Module'); r.bold = True; r.font.size = Pt(14)
    doc.add_paragraph('')
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run('This certifies that').font.size = Pt(12)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('[Learner Name]'); r.bold = True; r.font.size = Pt(24); r.font.color.rgb = RGBColor(0, 118, 112)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run('completed HIPAA in Practice, a HIPAA Privacy and Security self-paced learning module.').font.size = Pt(12)
    doc.add_paragraph('')
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run('HIPAA is the rule. Trust is the reason.').italic = True
    doc.add_paragraph('')
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run('Issued by RevExpert One  |  2026 Edition').font.size = Pt(10)
    doc.save(OUT / 'revexpert-one-certificate-template.docx')

def wrap(c, text, x, y, width, leading=14):
    words = text.split(); line = ''
    for word in words:
        test = (line + ' ' + word).strip()
        if stringWidth(test, 'Helvetica', 10) > width:
            c.drawString(x, y, line); y -= leading; line = word
        else: line = test
    if line: c.drawString(x, y, line); y -= leading
    return y

def make_pdf():
    path = OUT / 'critical-hipaa-terms.pdf'
    c = Canvas(str(path), pagesize=letter)
    width, height = letter
    c.setFillColor(Color(.04,.20,.22, alpha=.07)); c.setFont('Helvetica-Bold', 80)
    c.saveState(); c.translate(80, 330); c.rotate(35); c.drawString(0, 0, 'RevX1'); c.restoreState()
    c.setFillColor(HexColor('#0B3035')); c.setFont('Helvetica-Bold', 22); c.drawString(54, 736, 'Critical HIPAA Terms')
    c.setFont('Helvetica', 10); c.setFillColor(HexColor('#51636A')); c.drawString(54, 716, 'RevExpert One learning reference  |  Educational companion to Behind the Data')
    y = 682
    for heading, body in terms:
        if y < 105:
            c.showPage(); y = 736
            c.setFillColor(Color(.04,.20,.22, alpha=.07)); c.setFont('Helvetica-Bold', 80)
            c.saveState(); c.translate(80, 330); c.rotate(35); c.drawString(0, 0, 'RevX1'); c.restoreState()
        c.setFillColor(HexColor('#00766F')); c.setFont('Helvetica-Bold', 11); c.drawString(54, y, heading); y -= 16
        c.setFillColor(HexColor('#26343A')); c.setFont('Helvetica', 10); y = wrap(c, body, 54, y, 500); y -= 9
    c.setFont('Helvetica', 8); c.setFillColor(HexColor('#51636A'))
    c.drawString(54, 54, 'For training use. Consult your organization’s approved HIPAA policies and escalation procedures.')
    c.save()

make_docx(); make_pdf()
